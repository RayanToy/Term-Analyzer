# -*- coding: utf-8 -*-
"""Сходимость чисел между отчётами.

Проверяет, что:

* сумма «Количество» по всем книжным файлам совпадает с «Итого» сводного;
* лист «Матрица» по строке даёт то же «Итого»;
* листы по тематическим группам в сумме дают лист «Все единицы»;
* число строк «Найденные» = числу ненулевых строк «Все единицы».

Запуск:  python tools/check_totals.py
"""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import config


def read_rows(path: Path, sheet: str):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return []
    rows = [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if r and r[0]]
    wb.close()
    return rows


def main() -> int:
    problems = 0

    def check(condition: bool, message: str) -> None:
        nonlocal problems
        print(("  OK   " if condition else "  ОШИБКА ") + message)
        if not condition:
            problems += 1

    summary = config.OUT_SUMMARY_XLSX
    if not summary.exists():
        print(f"Нет сводного отчёта: {summary}")
        return 1

    print(f"Сводный отчёт: {summary.name}")

    # --- Итого из сводного -------------------------------------------------
    totals: dict[tuple[str, str], int] = {}
    for row in read_rows(summary, "Все единицы"):
        totals[(row[0], row[1])] = row[2] or 0

    # --- сумма по книжным файлам ------------------------------------------
    per_book: dict[tuple[str, str], int] = defaultdict(int)
    book_files = sorted(config.OUT_PER_BOOK.rglob("*_результат.xlsx"))
    print(f"Файлов по книгам: {len(book_files)}")
    for path in book_files:
        for row in read_rows(path, "Все единицы"):
            per_book[(row[0], row[1])] += row[2] or 0

    mismatched = [
        (key, per_book.get(key, 0), totals.get(key, 0))
        for key in set(totals) | set(per_book)
        if per_book.get(key, 0) != totals.get(key, 0)
    ]
    check(not mismatched, f"сумма по книгам == «Итого» сводного ({len(totals)} строк)")
    for key, a, b in mismatched[:5]:
        print(f"         {key}: по книгам {a}, в сводном {b}")

    # --- колонки-пометки ---------------------------------------------------
    # 3 — «в т.ч. из PDF», 4 — «в т.ч. добавлено вручную»
    for column, title in ((3, "в т.ч. из PDF"), (4, "в т.ч. добавлено вручную")):
        in_summary = {
            (row[0], row[1]): row[column] or 0
            for row in read_rows(summary, "Все единицы")
        }
        in_books: dict[tuple[str, str], int] = defaultdict(int)
        for path in book_files:
            for row in read_rows(path, "Все единицы"):
                in_books[(row[0], row[1])] += row[column] or 0
        bad = [
            key for key in set(in_summary) | set(in_books)
            if in_summary.get(key, 0) != in_books.get(key, 0)
        ]
        check(not bad, f"колонка «{title}» сходится (всего {sum(in_summary.values())})")

    # --- матрица -----------------------------------------------------------
    matrix_bad = []
    for row in read_rows(summary, "Матрица"):
        declared = row[2] or 0
        by_books = sum(v or 0 for v in row[3:])
        if declared != by_books:
            matrix_bad.append((row[0], declared, by_books))
    check(not matrix_bad, "строки «Матрицы» сходятся с колонкой «Итого»")
    for name, a, b in matrix_bad[:5]:
        print(f"         {name}: Итого {a}, по учебникам {b}")

    # --- группы ------------------------------------------------------------
    group_sum: dict[tuple[str, str], int] = {}
    wb = openpyxl.load_workbook(summary, data_only=True, read_only=True)
    group_sheets = [
        s for s in wb.sheetnames
        if s[0].isdigit() and "." in s[:3]
    ]
    wb.close()
    for sheet in group_sheets:
        for row in read_rows(summary, sheet):
            group_sum[(row[0], row[1])] = row[2] or 0
    check(
        group_sum == totals,
        f"листы по группам ({len(group_sheets)}) в сумме дают «Все единицы»",
    )

    # --- найденные ---------------------------------------------------------
    found = len(read_rows(summary, "Найденные"))
    nonzero = sum(1 for v in totals.values() if v)
    check(found == nonzero, f"«Найденные» = ненулевые строки «Все единицы» ({found})")

    print()
    print("Расхождений нет." if not problems else f"Расхождений: {problems}")
    return 0 if not problems else 1


if __name__ == "__main__":
    raise SystemExit(main())
